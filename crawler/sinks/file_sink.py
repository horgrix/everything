"""
Local file sink: exports cleaned rows to CSV / JSONL / JSON / Excel.

Claims the ``target_file`` output key.

Config::

    outputs:
      - target_file:
          path: "output/apps_{today}.csv"   # {today}/{now}/... template vars
          format: csv                       # csv | jsonl | json | excel
          mode: overwrite                   # overwrite | append
        parser: {...}

Rows are buffered and written on ``flush()`` (end of run), so ``overwrite``
mode produces one complete file per run and ``append`` accumulates across
runs.
"""

import asyncio
import csv
import json
import logging
import os

from .base import OutputSink, SinkResult
from ..template import URLTemplate

logger = logging.getLogger(__name__)

_FORMATS = {"csv", "jsonl", "json", "excel"}


def union_fields(rows: list[dict]) -> list[str]:
    """Column order = first-seen order across all rows."""
    fields: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                fields.append(key)
    return fields


def write_rows_to_file(rows: list[dict], path: str, fmt: str, mode: str) -> None:
    """Write ``rows`` to ``path`` (sync). Raises ValueError on bad format."""
    fmt = (fmt or "csv").lower()
    mode = (mode or "overwrite").lower()
    if fmt not in _FORMATS:
        raise ValueError(
            f"target_file.format 不支持: '{fmt}'（可选 {sorted(_FORMATS)}）"
        )

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)

    if fmt == "csv":
        _write_csv(rows, path, mode)
    elif fmt == "jsonl":
        _write_jsonl(rows, path, mode)
    elif fmt == "json":
        _write_json(rows, path, mode)
    else:
        _write_excel(rows, path, mode)


def _write_csv(rows: list[dict], path: str, mode: str) -> None:
    fields = union_fields(rows)
    append = mode == "append"
    write_header = not append or not os.path.exists(path) or os.path.getsize(path) == 0
    with open(path, "a" if append else "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        if write_header:
            writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k) for k in fields})


def _write_jsonl(rows: list[dict], path: str, mode: str) -> None:
    append = mode == "append"
    with open(path, "a" if append else "w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def _write_json(rows: list[dict], path: str, mode: str) -> None:
    data = list(rows)
    if mode == "append" and os.path.exists(path) and os.path.getsize(path) > 0:
        try:
            with open(path, "r", encoding="utf-8") as f:
                existing = json.load(f)
            if isinstance(existing, list):
                data = existing + data
        except (json.JSONDecodeError, OSError):
            logger.warning("target_file 追加模式：无法读取已有 JSON %s，将覆盖", path)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _write_excel(rows: list[dict], path: str, mode: str) -> None:
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "target_file format=excel requires openpyxl:\n"
            "  pip install openpyxl"
        )

    fields = union_fields(rows)
    if mode == "append" and os.path.exists(path):
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        has_header = ws.max_row >= 1 and ws.cell(row=1, column=1).value is not None
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        has_header = False

    if not has_header and fields:
        ws.append(fields)
    for row in rows:
        ws.append([row.get(k) for k in fields])

    wb.save(path)


class FileSink(OutputSink):
    """Local file export sink (key: ``target_file``)."""

    key = "target_file"

    def __init__(self):
        # id(output_config) -> (config_ref, rows)
        self._buffers: dict[int, tuple[dict, list[dict]]] = {}

    # ---- helpers ----

    def _buffer(self, output_config: dict) -> list[dict]:
        key = id(output_config)
        entry = self._buffers.get(key)
        if entry is None or entry[0] is not output_config:
            entry = (output_config, [])
            self._buffers[key] = entry
        return entry[1]

    @staticmethod
    def resolve_path(file_config: dict, context: dict) -> str:
        """Resolve template variables ({today}, {now}, ...) in the path."""
        return URLTemplate.resolve(file_config.get("path", ""), context=context or {})

    # ---- OutputSink interface ----

    def prepare(self, output_config: dict, context: dict, db) -> None:
        file_config = output_config.get("target_file")
        if not file_config:
            return
        self._buffer(output_config)
        path = self.resolve_path(file_config, context)
        if path:
            os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)

    def write(self, rows: list[dict], output_config: dict, context: dict, db) -> SinkResult:
        if not rows or not output_config.get("target_file"):
            return SinkResult()
        self._buffer(output_config).extend(rows)
        return SinkResult()

    def flush(self, output_config: dict, context: dict, db) -> SinkResult:
        file_config = output_config.get("target_file")
        if not file_config:
            return SinkResult()
        buffer = self._buffer(output_config)
        if not buffer:
            return SinkResult()
        rows = list(buffer)
        buffer.clear()

        path = self.resolve_path(file_config, context)
        if not path:
            logger.error("target_file 缺少 path，跳过写入")
            return SinkResult(failed=len(rows))

        write_rows_to_file(
            rows, path,
            file_config.get("format", "csv"),
            file_config.get("mode", "overwrite"),
        )
        logger.info("target_file 写入 %d 行到 %s", len(rows), path)
        # File is a non-DB destination: report as `sent` so `new`/`updated`
        # stay reserved for SQL sinks (avoids double counting when an output
        # writes to both target_table and target_file).
        return SinkResult(sent=len(rows))

    # ---- Async facade: file I/O off the event loop ----

    async def flush_async(self, output_config, context, db) -> SinkResult:
        return await asyncio.to_thread(self.flush, output_config, context, db)
